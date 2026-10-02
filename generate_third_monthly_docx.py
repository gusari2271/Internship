import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="B0C4DE", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def build_third_monthly_docx():
    doc = docx.Document()

    # Configure Margins: 0.75 in
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # =========================================================================
    # PAGE 1: ASSIGNMENT COVER SHEET
    # =========================================================================
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_top = p_top.add_run("Assignment No.: _____")
    r_top.font.name = "Calibri"
    r_top.font.size = Pt(9.5)

    p_uni = doc.add_paragraph()
    p_uni.paragraph_format.space_before = Pt(4)
    p_uni.paragraph_format.space_after = Pt(2)
    r_uni = p_uni.add_run("HELP University\n")
    r_uni.bold = True
    r_uni.font.name = "Calibri"
    r_uni.font.size = Pt(15)
    r_uni.font.color.rgb = RGBColor(192, 0, 0)
    r_uni_sub = p_uni.add_run("university of achievers")
    r_uni_sub.font.name = "Calibri"
    r_uni_sub.font.size = Pt(9)
    r_uni_sub.font.color.rgb = RGBColor(100, 100, 100)

    p_cover_title = doc.add_paragraph()
    p_cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cover_title.paragraph_format.space_before = Pt(10)
    p_cover_title.paragraph_format.space_after = Pt(12)
    r_ct = p_cover_title.add_run("Assignment Cover Sheet")
    r_ct.bold = True
    r_ct.font.name = "Calibri"
    r_ct.font.size = Pt(18)
    r_ct.font.color.rgb = RGBColor(20, 20, 20)

    # Student Info Table
    t_student = doc.add_table(rows=3, cols=2)
    t_student.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_student, color="808080", sz="4")
    t_student.rows[0].cells[0].width = Inches(5.2)
    t_student.rows[0].cells[1].width = Inches(1.8)

    # Row 0
    c00 = t_student.rows[0].cells[0]
    set_cell_background(c00, "E8EEF5")
    p = c00.paragraphs[0]
    r = p.add_run("Student Information (For group assignment, please state names of all members)")
    r.bold = True
    r.font.size = Pt(9)

    c01 = t_student.rows[0].cells[1]
    set_cell_background(c01, "E8EEF5")
    p = c01.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Grade/Marks")
    r.bold = True
    r.font.size = Pt(9)

    # Row 1: Header names
    c10 = t_student.rows[1].cells[0]
    p = c10.paragraphs[0]
    r = p.add_run("Name: ")
    r.bold = True
    r.font.size = Pt(9)
    r_val = p.add_run("I PUTU AGUS ARIBAWA")
    r_val.font.size = Pt(9)
    
    r_id = p.add_run("                                      ID: ")
    r_id.bold = True
    r_id.font.size = Pt(9)
    r_id_val = p.add_run("E2400080")
    r_id_val.font.size = Pt(9)

    c11 = t_student.rows[1].cells[1]
    c11.merge(t_student.rows[2].cells[1])

    # Row 2
    c20 = t_student.rows[2].cells[0]
    p = c20.paragraphs[0]
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(" ")

    p_sp1 = doc.add_paragraph()
    p_sp1.paragraph_format.space_before = Pt(4)
    p_sp1.paragraph_format.space_after = Pt(4)

    # Module Info Table
    t_module = doc.add_table(rows=7, cols=2)
    t_module.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_module, color="808080", sz="4")
    t_module.rows[0].cells[0].width = Inches(5.2)
    t_module.rows[0].cells[1].width = Inches(1.8)

    # Header Row
    c_m0 = t_module.rows[0].cells[0]
    set_cell_background(c_m0, "E8EEF5")
    p = c_m0.paragraphs[0]
    r = p.add_run("Module/Subject Information")
    r.bold = True
    r.font.size = Pt(9)

    c_m1 = t_module.rows[0].cells[1]
    set_cell_background(c_m1, "E8EEF5")
    p = c_m1.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Office Acknowledgement")
    r.bold = True
    r.font.size = Pt(9)

    # Merge right column for rows 1 to 6
    c_ack = t_module.rows[1].cells[1]
    for row_idx in range(2, 7):
        c_ack.merge(t_module.rows[row_idx].cells[1])

    mod_data = [
        ("Module/Subject Code", "BIT 320"),
        ("Module/Subject Name", "Industrial Internship"),
        ("Lecturer/Tutor/Facilitator", "Ms. Ayu Chrisniyanti S.Kom.,BIT.,MBA / Ms. Anitha Velayuktam"),
        ("Due Date", "2nd October 2026"),
        ("Assignment Title/Topic", "Third Monthly Progress Report and Logbook"),
        ("Intake / Word Count", "Sem July, 2026  |  Word Count: ~1,850 words"),
    ]

    for idx, (label, val) in enumerate(mod_data):
        cell = t_module.rows[1 + idx].cells[0]
        p = cell.paragraphs[0]
        r_lbl = p.add_run(f"{label}: ")
        r_lbl.bold = True
        r_lbl.font.size = Pt(8.5)
        r_v = p.add_run(val)
        r_v.font.size = Pt(8.5)

    p_dec = doc.add_paragraph()
    p_dec.paragraph_format.space_before = Pt(8)
    p_dec.paragraph_format.space_after = Pt(2)
    r_dec = p_dec.add_run("Declaration")
    r_dec.bold = True
    r_dec.font.size = Pt(9.5)

    bullets = [
        "I/We have read and understood the Programme Handbook that explains on plagiarism, and I/we testify that, unless otherwise acknowledged, the work submitted herein is entirely my/our own.",
        "I/We declare that no part of this assignment has been written for me/us by any other person(s) except where such collaboration has been authorized by the lecturer concerned.",
        "I/We authorize the University to test any work submitted by me/us, using text comparison software, for instances of plagiarism. I/We understand this will involve the University or its contractors copying my/our work and storing it on a database to be used in future to test work submitted by others."
    ]
    for b in bullets:
        pb = doc.add_paragraph()
        pb.paragraph_format.space_before = Pt(1)
        pb.paragraph_format.space_after = Pt(1)
        pb.paragraph_format.left_indent = Inches(0.15)
        r_b = pb.add_run("• ")
        r_b.bold = True
        r_b.font.size = Pt(8)
        r_txt = pb.add_run(b)
        r_txt.font.size = Pt(8)

    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(3)
    p_note.paragraph_format.space_after = Pt(4)
    r_n = p_note.add_run("Note: 1) The attachment of this statement on any electronically submitted assignments will be deemed to have the same authority as a signed statement.\n2) The Group Leader signs the declaration on behalf of all members.")
    r_n.font.size = Pt(7.5)
    r_n.font.italic = True

    # Signature Table
    t_sig = doc.add_table(rows=2, cols=2)
    t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_sig, color="808080", sz="4")
    t_sig.rows[0].cells[0].width = Inches(4.5)
    t_sig.rows[0].cells[1].width = Inches(2.5)

    p_s0 = t_sig.rows[0].cells[0].paragraphs[0]
    r = p_s0.add_run("Signature: ")
    r.bold = True
    r.font.size = Pt(8.5)
    r_sig = p_s0.add_run("[DIGITAL SIGNATURE / I PUTU AGUS ARIBAWA]")
    r_sig.font.size = Pt(8.5)

    p_s1 = t_sig.rows[0].cells[1].paragraphs[0]
    r = p_s1.add_run("Date: ")
    r.bold = True
    r.font.size = Pt(8.5)
    r_dt = p_s1.add_run("2 October 2026")
    r_dt.font.size = Pt(8.5)

    # Row 1 merged: Email
    c_em = t_sig.rows[1].cells[0]
    c_em.merge(t_sig.rows[1].cells[1])
    p_em = c_em.paragraphs[0]
    r = p_em.add_run("E-mail: ")
    r.bold = True
    r.font.size = Pt(8.5)
    r_em = p_em.add_run("230030591@stikom-bali.ac.id")
    r_em.font.size = Pt(8.5)

    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(8)
    p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_ft = p_foot.add_run("ACA-F-020(010611:01)  |  Page 1 of 2")
    r_ft.font.size = Pt(7.5)
    r_ft.font.color.rgb = RGBColor(120, 120, 120)

    # =========================================================================
    # PAGE 2: FEEDBACK / COMMENTS SHEET
    # =========================================================================
    doc.add_page_break()

    p_fb_title = doc.add_paragraph()
    p_fb_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fb_title.paragraph_format.space_before = Pt(6)
    p_fb_title.paragraph_format.space_after = Pt(10)
    r_fbt = p_fb_title.add_run("Feedback / Comments Sheet")
    r_fbt.bold = True
    r_fbt.font.name = "Calibri"
    r_fbt.font.size = Pt(16)

    t_fb = doc.add_table(rows=4, cols=1)
    t_fb.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_fb, color="808080", sz="4")
    t_fb.rows[0].cells[0].width = Inches(7.0)

    c_fb0 = t_fb.rows[0].cells[0]
    set_cell_background(c_fb0, "E8EEF5")
    p = c_fb0.paragraphs[0]
    r = p.add_run("Feedback / Comments*")
    r.bold = True
    r.font.size = Pt(9.5)

    fb_sections = [
        ("Main Strengths", 7),
        ("Main Weaknesses", 5),
        ("Suggestions for improvement", 5)
    ]
    for idx, (sec_title, lines_cnt) in enumerate(fb_sections):
        cell = t_fb.rows[1 + idx].cells[0]
        set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
        p = cell.paragraphs[0]
        r = p.add_run(sec_title)
        r.bold = True
        r.font.size = Pt(9)
        for _ in range(lines_cnt):
            p_sub = cell.add_paragraph()
            p_sub.paragraph_format.space_before = Pt(4)
            p_sub.paragraph_format.space_after = Pt(4)
            p_sub.add_run(" ")

    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(10)
    p_sp2.paragraph_format.space_after = Pt(4)

    t_ack = doc.add_table(rows=3, cols=2)
    t_ack.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_ack, color="808080", sz="4")
    t_ack.rows[0].cells[0].width = Inches(3.5)
    t_ack.rows[0].cells[1].width = Inches(3.5)

    c_ack0 = t_ack.rows[0].cells[0]
    c_ack0.merge(t_ack.rows[0].cells[1])
    set_cell_background(c_ack0, "E8EEF5")
    p = c_ack0.paragraphs[0]
    r = p.add_run("Student acknowledge feedback/comments")
    r.bold = True
    r.font.size = Pt(9)

    p0 = t_ack.rows[1].cells[0].paragraphs[0]
    p0.paragraph_format.space_before = Pt(4)
    p0.paragraph_format.space_after = Pt(20)
    r = p0.add_run("Grader's signature: ______________________")
    r.bold = True
    r.font.size = Pt(8.5)

    p1 = t_ack.rows[1].cells[1].paragraphs[0]
    p1.paragraph_format.space_before = Pt(4)
    p1.paragraph_format.space_after = Pt(20)
    r = p1.add_run("Student's signature: [I PUTU AGUS ARIBAWA]")
    r.bold = True
    r.font.size = Pt(8.5)

    p2 = t_ack.rows[2].cells[0].paragraphs[0]
    r = p2.add_run("Date: ______________________")
    r.bold = True
    r.font.size = Pt(8.5)

    p3 = t_ack.rows[2].cells[1].paragraphs[0]
    r = p3.add_run("Date: 2 October 2026")
    r.bold = True
    r.font.size = Pt(8.5)

    p_fb_note = doc.add_paragraph()
    p_fb_note.paragraph_format.space_before = Pt(8)
    r = p_fb_note.add_run(
        "Note: 1) A soft and hard copy of the assignment shall be submitted.\n"
        "2) The signed copy of the assignment cover sheet shall be retained by the marker.\n"
        "3) If the Turnitin report is required, students have to submit it with the assignment. "
        "However, departments may allow students up to THREE (3) working days after submission to submit the Turnitin report.\n"
        "*Use additional sheets if required."
    )
    r.font.size = Pt(7.5)
    r.font.italic = True

    p_foot2 = doc.add_paragraph()
    p_foot2.paragraph_format.space_before = Pt(6)
    p_foot2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_ft2 = p_foot2.add_run("ACA-F-020(010611:01)  |  Page 2 of 2")
    r_ft2.font.size = Pt(7.5)
    r_ft2.font.color.rgb = RGBColor(120, 120, 120)

    # =========================================================================
    # PAGE 3 ONWARDS: MONTHLY PROGRESS REPORT (APPENDIX C)
    # =========================================================================
    doc.add_page_break()

    p_app_hdr = doc.add_paragraph()
    p_app_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_app = p_app_hdr.add_run("BIT320 Industrial Internship Semester 2, 2026\nAppendix C")
    r_app.bold = True
    r_app.font.name = "Calibri"
    r_app.font.size = Pt(9.5)
    r_app.font.color.rgb = RGBColor(60, 60, 60)

    p_mpr_title = doc.add_paragraph()
    p_mpr_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mpr_title.paragraph_format.space_before = Pt(6)
    p_mpr_title.paragraph_format.space_after = Pt(12)
    r_mpr = p_mpr_title.add_run("Monthly Progress Report")
    r_mpr.bold = True
    r_mpr.font.name = "Calibri"
    r_mpr.font.size = Pt(16)
    r_mpr.font.color.rgb = RGBColor(20, 20, 20)

    # Table for Appendix C Form
    # Rows:
    # 0: Project Name
    # 1: Student Name and ID
    # 2: Date
    # 3: Reporting Period
    # 4: Work completed this reporting period
    # 5: Work to complete next reporting period
    # 6: What is going well and why
    # 7: What is not going well and why
    # 8: Suggestions/Issues
    # 9: Changes to internship terms
    t_mpr = doc.add_table(rows=10, cols=1)
    t_mpr.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_mpr, color="808080", sz="4")
    t_mpr.rows[0].cells[0].width = Inches(7.0)

    # Helper function to populate cell
    def populate_header_field(cell, label, value):
        set_cell_margins(cell, top=60, bottom=60, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r_lbl = p.add_run(label)
        r_lbl.bold = True
        r_lbl.font.name = "Calibri"
        r_lbl.font.size = Pt(9.5)
        r_v = p.add_run(f" {value}")
        r_v.font.name = "Calibri"
        r_v.font.size = Pt(9.5)

    populate_header_field(t_mpr.rows[0].cells[0], "Project Name:", "Grahita Design Website")
    populate_header_field(t_mpr.rows[1].cells[0], "Student Name and ID:", "I Putu Agus Aribawa (E2400080)")
    populate_header_field(t_mpr.rows[2].cells[0], "Date:", "2 October 2026")
    populate_header_field(t_mpr.rows[3].cells[0], "Reporting Period:", "2 September – 2 October 2026")

    # Section 4: Work completed this reporting period
    c_wc = t_mpr.rows[4].cells[0]
    set_cell_margins(c_wc, top=100, bottom=100, left=120, right=120)
    p_wc_title = c_wc.paragraphs[0]
    p_wc_title.paragraph_format.space_before = Pt(3)
    p_wc_title.paragraph_format.space_after = Pt(4)
    r = p_wc_title.add_run("Work completed this reporting period:")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(10)

    completed_bullets = [
        "Completed Administrative Authentication & Access Control: Implemented the complete AuthModule and AuthService in NestJS with bcrypt password hashing (10 salt rounds) and JWT signing. Created the User entity to represent administrator credentials, built the POST /auth/login endpoint, and applied the JwtAuthGuard across all mutation endpoints (POST /projects, PUT /projects/:id, DELETE /projects/:id) in ProjectsController to restrict administrative actions while keeping public portfolio read routes accessible.",
        "Finalized Bilingual API & Localization Pipeline: Completed the bilingual architecture for portfolio projects and contact inquiries (title_id/title_en, description_id/description_en). Standardized query parameter handling (?lang=id or ?lang=en) and verified consistent response formatting with Kevin for the Angular frontend language toggle.",
        "Implemented Global Exception Filter & Pagination: Built a NestJS global exception filter to standardize all error responses into a consistent JSON envelope with timestamps, status codes, and user-friendly error messages. Implemented database pagination (page and limit with TypeORM skip and take) and category filtering on GET /projects, verifying boundary edge cases via Postman.",
        "API Documentation & Seeder Updates: Configured @nestjs/swagger in main.ts with comprehensive OpenAPI decorators (@ApiTags, @ApiOperation, @ApiResponse), generating interactive API documentation at /api/docs. Enriched the database seeder (onModuleInit) with realistic bilingual descriptions and verified gallery image datasets for all default architectural projects.",
        "Automated Unit Testing with Jest: Authored automated unit test suites for ProjectsService and AuthService using Jest, testing CRUD methods, credential verification, and cube index conflict resolution with 100% passing assertions.",
        "Production Deployment Preparation & Server Configuration: Formulated the deployment roadmap, disabled TypeORM's automatic synchronize: true in favor of production migration safety, and created .env.production. Provisioned cloud hosting with Node.js LTS, PM2 process management with auto-restart, and Nginx reverse proxy.",
        "Live Deployment & Full-Stack Integration: Deployed the NestJS backend and SQLite database to production, configured production CORS to whitelist Kevin's live Angular domain, and resolved production HTTPS image URL generation. Conducted comprehensive end-to-end smoke testing across all public and protected routes.",
        "Security Hardening: Implemented API rate limiting using @nestjs/throttler (ThrottlerModule), applying strict rate limits on /contact and /auth/login to prevent spam and brute-force attacks. Configured helmet security middleware to enforce HTTP security headers against XSS, clickjacking, and MIME sniffing.",
        "Formal User Acceptance Testing (UAT): Conducted a formal UAT session on 23 September 2026 with client Bli Dek, Kevin, and the supervisor. Validated project browsing, 3D WebGL spatial cube slot navigation, client inquiry submissions, and admin dashboard updates, securing formal client UAT sign-off.",
        "Client Feedback Adjustments & Database Backups: Refined database records with high-resolution photography and revised architectural project descriptions provided by Bli Dek. Developed an automated SQLite snapshot backup script and validated database disaster recovery procedures.",
        "Automated Email Notification Pipeline & Bug Fixing: Configured the MailModule and Nodemailer transporter in ContactService to automatically dispatch new inquiry notifications to studio administrators and delivery confirmation receipts to prospective clients. Collaborated with Kevin to resolve table vertical alignment and action button positioning in the admin dashboard catalog view.",
        "Formal Client Handover & Project Finalization: Conducted the official project handover meeting with Bli Dek on 30 September 2026, delivering administrative credentials, system documentation, and receiving enthusiastic client acceptance. Finalized README.md, authored the BIT320 Final Report (FINAL_REPORT_BIT320.md), and attended the supervisor exit evaluation meeting on 2 October 2026.",
        "Weekly Educational Content Creation: Researched, designed, and published four educational Instagram carousel posts on mobile phone fun facts (display technologies, camera sensors, battery/charging, and mobile processors) every Tuesday, in addition to recording and publishing the monthly educational video for September."
    ]

    for b in completed_bullets:
        pb = c_wc.add_paragraph()
        pb.paragraph_format.space_before = Pt(2)
        pb.paragraph_format.space_after = Pt(3)
        pb.paragraph_format.left_indent = Inches(0.18)
        r_b = pb.add_run("● ")
        r_b.bold = True
        r_b.font.size = Pt(9)
        r_b.font.color.rgb = RGBColor(0, 32, 96)
        r_txt = pb.add_run(b)
        r_txt.font.name = "Calibri"
        r_txt.font.size = Pt(9)
        r_txt.font.color.rgb = RGBColor(30, 30, 30)

    # Section 5: Work to complete next reporting period
    c_wn = t_mpr.rows[5].cells[0]
    set_cell_margins(c_wn, top=100, bottom=100, left=120, right=120)
    p_wn_title = c_wn.paragraphs[0]
    p_wn_title.paragraph_format.space_before = Pt(3)
    p_wn_title.paragraph_format.space_after = Pt(4)
    r = p_wn_title.add_run("Work to complete next reporting period:")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(10)

    next_bullets = [
        "Internship Milestone Completion: This represents the third and final monthly progress report, marking the formal conclusion of the 14-week BIT320 Industrial Internship. All assigned backend engineering deliverables, full-stack integrations, client handover requirements, and academic deliverables have been 100% completed and officially submitted.",
        "Post-Internship Maintenance & Operational Handoff: Provide ongoing advisory support during the initial weeks of independent studio operation by GRAHITA Design personnel, monitoring cloud server uptime, memory consumption, and SQLite database storage growth as portfolio projects and client inquiries accumulate over the upcoming quarters.",
        "Long-Term Architectural Enhancements: Hand over technical documentation for potential future architectural upgrades, including implementing role-based access control (RBAC) if multiple administrative tiers are onboarded, and migrating gallery uploads to external cloud object storage (e.g. AWS S3 or Cloudinary) if studio photography exceeds local VPS storage limits."
    ]

    for b in next_bullets:
        pb = c_wn.add_paragraph()
        pb.paragraph_format.space_before = Pt(2)
        pb.paragraph_format.space_after = Pt(3)
        pb.paragraph_format.left_indent = Inches(0.18)
        r_b = pb.add_run("● ")
        r_b.bold = True
        r_b.font.size = Pt(9)
        r_b.font.color.rgb = RGBColor(0, 32, 96)
        r_txt = pb.add_run(b)
        r_txt.font.name = "Calibri"
        r_txt.font.size = Pt(9)
        r_txt.font.color.rgb = RGBColor(30, 30, 30)

    # Section 6: What is going well and why
    c_gw = t_mpr.rows[6].cells[0]
    set_cell_margins(c_gw, top=100, bottom=100, left=120, right=120)
    p_gw_title = c_gw.paragraphs[0]
    p_gw_title.paragraph_format.space_before = Pt(3)
    p_gw_title.paragraph_format.space_after = Pt(4)
    r = p_gw_title.add_run("What is going well and why:")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(10)

    p_gw_body = c_gw.add_paragraph()
    p_gw_body.paragraph_format.space_before = Pt(2)
    p_gw_body.paragraph_format.space_after = Pt(4)
    r_gw = p_gw_body.add_run(
        "Backend development progressed with high velocity and architectural stability throughout this final month. "
        "The modular Controller-Service-Repository architecture of NestJS and the relational SQLite database designed during "
        "earlier weeks provided an exceptionally solid foundation. Adding advanced features—such as JWT route guards, "
        "@nestjs/throttler rate limiting, helmet security headers, and the MailModule email pipeline—was seamless and did not "
        "destabilize existing modules. Deploying the backend to the live cloud staging environment early in Week 12 gave the team "
        "ample buffer time to conduct smoke tests, preview features with the client, and execute a formal UAT in Week 13 without "
        "last-minute deadline pressure. Furthermore, daily coordination with frontend developer Kevin Wiratama remained highly "
        "effective; transitioning from verbal discussions to interactive Swagger/OpenAPI documentation and synchronized Postman "
        "collections eliminated API contract discrepancies and accelerated frontend-backend integration. The custom "
        "handleCubeIndexConflict algorithm worked reliably during client demonstrations, ensuring 3D spatial cube navigation "
        "was robust. Finally, the formal handover meeting with Bli Dek was an overwhelming success, with the client expressing high "
        "satisfaction regarding website responsiveness, bilingual capabilities, and ease of portfolio content management."
    )
    r_gw.font.name = "Calibri"
    r_gw.font.size = Pt(9)
    r_gw.font.color.rgb = RGBColor(30, 30, 30)

    # Section 7: What is not going well and why
    c_nw = t_mpr.rows[7].cells[0]
    set_cell_margins(c_nw, top=100, bottom=100, left=120, right=120)
    p_nw_title = c_nw.paragraphs[0]
    p_nw_title.paragraph_format.space_before = Pt(3)
    p_nw_title.paragraph_format.space_after = Pt(4)
    r = p_nw_title.add_run("What is not going well and why:")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(10)

    p_nw_body = c_nw.add_paragraph()
    p_nw_body.paragraph_format.space_before = Pt(2)
    p_nw_body.paragraph_format.space_after = Pt(4)
    r_nw = p_nw_body.add_run(
        "The primary challenge encountered this month stemmed from environmental differences between the local development environment "
        "and cloud production hosting. When deploying the application, uploaded image URLs generated in projects.controller.ts initially "
        "defaulted to localhost rather than the production HTTPS domain, requiring dynamic request protocol and host header resolution to ensure "
        "uploaded photos loaded properly. Additionally, configuring Nodemailer in production surfaced warning logs when SMTP environment variables "
        "were missing or misconfigured; this was addressed by implementing graceful error handling and fallback console logging in MailService so "
        "contact inquiries safely persisted to the database regardless of mail transport status. Finally, handling complex updates on the Admin Dashboard "
        "that combined project text metadata, gallery image arrays, and 3D cube index assignments required meticulous TypeORM entity relationship management "
        "to prevent child image detachment, which was systematically resolved by implementing explicit image retention tracking (keepImageIds)."
    )
    r_nw.font.name = "Calibri"
    r_nw.font.size = Pt(9)
    r_nw.font.color.rgb = RGBColor(30, 30, 30)

    # Section 8: Suggestions/Issues
    c_sug = t_mpr.rows[8].cells[0]
    set_cell_margins(c_sug, top=100, bottom=100, left=120, right=120)
    p_sug_title = c_sug.paragraphs[0]
    p_sug_title.paragraph_format.space_before = Pt(3)
    p_sug_title.paragraph_format.space_after = Pt(4)
    r = p_sug_title.add_run("Suggestions/Issues:")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(10)

    p_sug_body = c_sug.add_paragraph()
    p_sug_body.paragraph_format.space_before = Pt(2)
    p_sug_body.paragraph_format.space_after = Pt(4)
    r_sug = p_sug_body.add_run(
        "Incorporating automated API documentation tools such as @nestjs/swagger from the very first week of multi-developer projects "
        "significantly streamlines cross-team integration by establishing an authoritative, interactive contract for all payload shapes. "
        "Furthermore, provisioning a production-like cloud staging environment early in the project lifecycle allows deployment-specific "
        "nuances (such as HTTPS domain resolution, CORS whitelisting, and filesystem permissions) to be surfaced and eliminated well ahead "
        "of client demonstrations. Finally, establishing automated database backup routines should be treated as a baseline practice immediately "
        "upon deploying to staging, ensuring that live client testing and administrative updates can proceed with complete disaster recovery safety."
    )
    r_sug.font.name = "Calibri"
    r_sug.font.size = Pt(9)
    r_sug.font.color.rgb = RGBColor(30, 30, 30)

    # Section 9: Changes to internship terms
    c_chg = t_mpr.rows[9].cells[0]
    set_cell_margins(c_chg, top=100, bottom=100, left=120, right=120)
    p_chg_title = c_chg.paragraphs[0]
    p_chg_title.paragraph_format.space_before = Pt(3)
    p_chg_title.paragraph_format.space_after = Pt(4)
    r = p_chg_title.add_run("Changes to internship terms (if applicable - job scope, location, working hours, supervisor):")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(10)

    p_chg_body = c_chg.add_paragraph()
    p_chg_body.paragraph_format.space_before = Pt(2)
    p_chg_body.paragraph_format.space_after = Pt(6)
    r_chg = p_chg_body.add_run(
        "None. All internship terms, job scope (Backend Developer), workplace location at Code Cipta, working hours, and supervisory "
        "arrangements remained consistent throughout the entire 14-week placement."
    )
    r_chg.font.name = "Calibri"
    r_chg.font.size = Pt(9)
    r_chg.font.color.rgb = RGBColor(30, 30, 30)

    # Signatures below table
    p_sig_bottom = doc.add_paragraph()
    p_sig_bottom.paragraph_format.space_before = Pt(14)
    p_sig_bottom.paragraph_format.space_after = Pt(4)
    r_sb1 = p_sig_bottom.add_run("Student's Signature : ")
    r_sb1.bold = True
    r_sb1.font.size = Pt(9.5)
    r_sb1_val = p_sig_bottom.add_run("[DIGITAL SIGNATURE / I PUTU AGUS ARIBAWA]")
    r_sb1_val.font.size = Pt(9.5)

    p_sig_bottom2 = doc.add_paragraph()
    p_sig_bottom2.paragraph_format.space_before = Pt(4)
    p_sig_bottom2.paragraph_format.space_after = Pt(4)
    r_name = p_sig_bottom2.add_run("Name : I Putu Agus Aribawa                                              Date: 2 October 2026")
    r_name.bold = True
    r_name.font.size = Pt(9.5)

    output_path = "c:\\Users\\LENOVO\\Downloads\\CLONE GRAHITA\\Internship\\third_monthly_report.docx"
    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

    # Also save a copy as 'third monthly.docx' for convenience
    copy_path = "c:\\Users\\LENOVO\\Downloads\\CLONE GRAHITA\\Internship\\third monthly.docx"
    doc.save(copy_path)
    print(f"Successfully generated copy: {copy_path}")

if __name__ == "__main__":
    build_third_monthly_docx()
