import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def build_logbook_docx():
    doc = docx.Document()

    # Set margins: 0.75 in
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        # Header setup
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("BIT320 Industrial Internship — Appendix B")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(120, 120, 120)

    # Document Header Title
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p_meta.add_run("HELP UNIVERSITY — BIT320 INDUSTRIAL INTERNSHIP\n")
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(192, 0, 0) # University red accent

    r2 = p_meta.add_run("IT INTERNSHIP LOGBOOK (Weeks 10 – 14)\n")
    r2.bold = True
    r2.font.name = "Calibri"
    r2.font.size = Pt(14)
    r2.font.color.rgb = RGBColor(30, 30, 30)

    r3 = p_meta.add_run(
        "Student Name: I Putu Agus Aribawa  |  Student ID: E2400080  |  Programme: BIT (Hons)\n"
        "Host Organisation: Code Cipta  |  Client: GRAHITA Design  |  Role: Backend Developer\n"
        "Period: 2 September 2026 – 2 October 2026"
    )
    r3.font.name = "Calibri"
    r3.font.size = Pt(9.5)
    r3.font.color.rgb = RGBColor(90, 90, 90)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(12)

    weeks_data = [
        {
            "week_title": "Week 10 (Continued)",
            "entries": [
                (
                    "Wednesday,\n2 September 2026",
                    [
                        "Implemented the AuthModule and AuthService in NestJS to establish the core authentication infrastructure for administrative access.",
                        "Created the User entity to represent administrator credentials, incorporating secure password hashing via bcrypt with appropriate salt rounds.",
                        "Built the POST /auth/login endpoint to validate administrator credentials against stored hashes and issue signed JSON Web Tokens (JWT) containing administrative claims."
                    ],
                    [
                        "JWT Authentication Architecture",
                        "Password Hashing with Bcrypt",
                        "NestJS AuthModule Configuration"
                    ]
                ),
                (
                    "Thursday,\n3 September 2026",
                    [
                        "Applied the JwtAuthGuard to mutation endpoints (POST /projects, PUT /projects/:id, DELETE /projects/:id) in ProjectsController to restrict content modifications to authenticated administrators.",
                        "Verified that public read endpoints (GET /projects, GET /projects/:id, GET /projects/pane-status) remain accessible without authentication for portfolio visitors.",
                        "Conducted comprehensive Postman tests using valid Bearer tokens, expired tokens, and missing authorization headers to verify proper HTTP 401 Unauthorized handling."
                    ],
                    [
                        "Route Protection with Guards",
                        "Postman Authorization Header Testing",
                        "HTTP Security Status Verification"
                    ]
                ),
                (
                    "Friday,\n4 September 2026",
                    [
                        "Committed and pushed the completed authentication module, guards, and entity files to the GitHub repository with clean, descriptive commit messages.",
                        "Held a weekly progress review meeting with the supervisor to demonstrate the completed JWT authentication flow and route protection on administrative endpoints.",
                        "Received positive supervisor feedback on the security architecture and outlined next week's focus on refining the bilingual API response structure, pagination, and error filtering."
                    ],
                    [
                        "Git Commit Documentation Skills",
                        "Technical Demonstration Skills",
                        "Stakeholder Feedback Incorporation"
                    ]
                )
            ]
        },
        {
            "week_title": "Week 11",
            "entries": [
                (
                    "Monday,\n7 September 2026",
                    [
                        "Conducted weekly sprint planning to break down technical priorities, focusing on finalizing the bilingual API response structure, pagination, and global error handling.",
                        "Applied time-blocking techniques to balance backend coding tasks, social media content creation, and coordination with Kevin for frontend integration.",
                        "Reviewed bilingual content requirements across the Project entity (title_id, title_en, description_id, description_en) to ensure data consistency between Indonesian and English versions."
                    ],
                    [
                        "Sprint Task Prioritization",
                        "Time-Blocking Technique",
                        "Bilingual Data Modeling"
                    ]
                ),
                (
                    "Tuesday,\n8 September 2026",
                    [
                        "Researched, designed, and published this week's educational Instagram carousel post about mobile phone fun facts, focusing on the evolution of mobile display and touchscreen technologies (from resistive panels to capacitive AMOLED).",
                        "Recorded, edited, and published the monthly educational video for September, explaining how smartphone miniaturization and processor fabrication nodes have evolved over the decades.",
                        "Collaborated with Kevin to standardize the bilingual API response shape, confirming query parameter handling (?lang=id or ?lang=en) for the frontend language switcher."
                    ],
                    [
                        "Educational Video Production",
                        "Social Media Content Strategy",
                        "API Contract Standardization"
                    ]
                ),
                (
                    "Wednesday,\n9 September 2026",
                    [
                        "Implemented pagination (page and limit query parameters with TypeORM skip and take) and category filtering on the GET /projects endpoint in ProjectsService.",
                        "Implemented a global exception filter in NestJS to standardize error responses into a consistent JSON envelope across the entire backend application.",
                        "Validated pagination boundary edge cases in Postman, verifying that requests with out-of-range offsets return appropriate empty arrays without throwing unhandled server exceptions."
                    ],
                    [
                        "Database Pagination with TypeORM",
                        "NestJS Global Exception Filters",
                        "Edge-Case Query Testing"
                    ]
                ),
                (
                    "Thursday,\n10 September 2026",
                    [
                        "Configured @nestjs/swagger in main.ts and decorated controllers and DTOs with OpenAPI annotations (@ApiTags, @ApiOperation, @ApiResponse) to produce interactive API documentation.",
                        "Updated the automated database seeder (onModuleInit) with enriched bilingual descriptions and verified gallery image arrays for all default architectural projects.",
                        "Shared the generated Swagger UI endpoint (/api/docs) and updated Postman collection with Kevin to facilitate seamless frontend integration."
                    ],
                    [
                        "Swagger/OpenAPI Documentation",
                        "Database Seeder Enrichment",
                        "Developer Tooling Collaboration"
                    ]
                ),
                (
                    "Friday,\n11 September 2026",
                    [
                        "Authored unit tests using Jest for ProjectsService and AuthService, testing CRUD methods, credential verification, and cube index conflict handling.",
                        "Conducted a weekly progress review and code demonstration with the supervisor, showcasing the live Swagger documentation and passing Jest test suites.",
                        "Discussed deployment prerequisites and hosting options with the supervisor, compiling a preparation checklist for next week's staging release."
                    ],
                    [
                        "Unit Testing with Jest",
                        "Test-Driven Verification Skills",
                        "Production Readiness Assessment"
                    ]
                )
            ]
        },
        {
            "week_title": "Week 12",
            "entries": [
                (
                    "Monday,\n14 September 2026",
                    [
                        "Formulated a comprehensive deployment roadmap for the backend, identifying necessary server environment configurations and security requirements.",
                        "Disabled TypeORM's automatic synchronize: true for the production configuration, migrating to safer explicit schema synchronization and backup procedures.",
                        "Set up environment variable management with .env.production, separating development database paths and secrets from production settings."
                    ],
                    [
                        "Production Deployment Planning",
                        "Database Migration Safety Practices",
                        "Environment Variable Segregation"
                    ]
                ),
                (
                    "Tuesday,\n15 September 2026",
                    [
                        "Researched, designed, and published this week's educational Instagram post covering mobile phone fun facts, highlighting the history of camera sensors from the first 0.11 MP phone camera in 2000 to modern periscope zoom lenses.",
                        "Implemented application-level request logging using NestJS built-in Logger to track incoming HTTP requests, response status codes, and execution latency.",
                        "Configured production static asset serving in main.ts, ensuring uploaded portfolio images in /uploads are served with proper caching headers and security checks."
                    ],
                    [
                        "Mobile Tech History Research",
                        "NestJS Built-in Logger Implementation",
                        "Static Asset Delivery Optimization"
                    ]
                ),
                (
                    "Wednesday,\n16 September 2026",
                    [
                        "Set up the production cloud hosting environment, installing Node.js LTS, PM2 process manager, and Nginx as a reverse proxy.",
                        "Configured PM2 ecosystem files to keep the NestJS application running continuously with automatic restarts upon system reboots or runtime failures.",
                        "Deployed the NestJS backend application to the live server, verified that the SQLite database file initialized correctly, and confirmed the health-check route."
                    ],
                    [
                        "Cloud Server Provisioning",
                        "PM2 Process Management",
                        "Nginx Reverse Proxy Configuration"
                    ]
                ),
                (
                    "Thursday,\n17 September 2026",
                    [
                        "Connected the deployed backend with Kevin's deployed frontend, updating CORS settings in main.ts to whitelist the live frontend production domain.",
                        "Attended a scheduled client demonstration meeting with Bli Dek, the supervisor, and Kevin, previewing the live staging website, portfolio layout, and 3D spatial cube navigation.",
                        "Gathered preliminary feedback from Bli Dek regarding image upload speed and specific architectural terminology on project detail pages."
                    ],
                    [
                        "Production CORS Configuration",
                        "Client Staging Demonstration",
                        "Client Feedback Gathering"
                    ]
                ),
                (
                    "Friday,\n18 September 2026",
                    [
                        "Resolved a minor URL resolution issue in projects.controller.ts where image paths needed to reflect the production HTTPS domain rather than localhost.",
                        "Ran end-to-end integration tests between the live frontend and backend alongside Kevin, confirming image uploads and contact form submissions operate reliably.",
                        "Held a weekly progress review meeting with the supervisor to assess the deployed staging environment and plan next week's formal User Acceptance Testing (UAT)."
                    ],
                    [
                        "Production Domain URL Resolution",
                        "End-to-End System Smoke Testing",
                        "Weekly Deployment Review"
                    ]
                )
            ]
        },
        {
            "week_title": "Week 13",
            "entries": [
                (
                    "Monday,\n21 September 2026",
                    [
                        "Outlined the sprint goals for the week, prioritizing client feedback items from Bli Dek, rate limiting, and drafting the BIT320 Final Report.",
                        "Applied the Eisenhower Matrix to organize tasks into urgent client adjustments versus non-urgent report documentation tasks to prevent schedule slippage.",
                        "Created a detailed User Acceptance Testing (UAT) test scenario checklist covering all core user journeys for client verification."
                    ],
                    [
                        "Eisenhower Matrix Prioritization",
                        "UAT Scenario Design",
                        "Release Schedule Management"
                    ]
                ),
                (
                    "Tuesday,\n22 September 2026",
                    [
                        "Researched, drafted, and posted this week's educational Instagram carousel on mobile phone fun facts, discussing the evolution of smartphone battery technology and fast charging (from NiCad to Silicon-Carbon and GaN chargers).",
                        "Implemented rate limiting across the API using @nestjs/throttler (ThrottlerModule), applying strict limits on the POST /contact and POST /auth/login endpoints to prevent spam and brute-force attacks.",
                        "Added helmet security middleware to HTTP response headers to protect against common web vulnerabilities like cross-site scripting (XSS) and clickjacking."
                    ],
                    [
                        "Educational Tech Storytelling",
                        "API Rate Limiting with Throttler",
                        "HTTP Security Header Hardening"
                    ]
                ),
                (
                    "Wednesday,\n23 September 2026",
                    [
                        "Conducted the formal User Acceptance Testing (UAT) session with client Bli Dek, Kevin, and the supervisor, systematically validating each feature against the project requirements.",
                        "Guided the client through submitting architectural inquiries via the contact form, verifying project filters, and testing administrative updates in the dashboard.",
                        "Documented the client's official UAT approval notes and noted minor requests for updated architectural project descriptions and image sequence adjustments."
                    ],
                    [
                        "User Acceptance Testing Protocol",
                        "Client-Facing Feature Walkthrough",
                        "Technical Defect and Request Logging"
                    ]
                ),
                (
                    "Thursday,\n24 September 2026",
                    [
                        "Updated project database records with the revised architectural project descriptions, locations, and high-resolution cover photos provided by Bli Dek.",
                        "Authored an automated SQLite database backup script that creates scheduled timestamped snapshots of db.sqlite to protect against accidental data loss.",
                        "Verified database restore functionality by successfully restoring a test backup copy without data corruption or index mismatches."
                    ],
                    [
                        "Production Data Content Refinement",
                        "SQLite Automated Backup Scripting",
                        "Database Disaster Recovery Testing"
                    ]
                ),
                (
                    "Friday,\n25 September 2026",
                    [
                        "Conducted API performance benchmarking using autocannon, confirming that the NestJS backend handles concurrent portfolio queries with sub-50ms latency.",
                        "Began drafting the initial sections of the BIT320 Industrial Internship Final Report, focusing on project background, system architecture, and SDLC methodology.",
                        "Met with the supervisor for the weekly progress review, presenting the successful UAT sign-off document and discussing the outline for the Final Report."
                    ],
                    [
                        "API Performance Benchmarking",
                        "Technical Academic Report Drafting",
                        "Supervisory Milestone Alignment"
                    ]
                )
            ]
        },
        {
            "week_title": "Week 14",
            "entries": [
                (
                    "Monday,\n28 September 2026",
                    [
                        "Conducted the final sprint planning session of the internship, creating a comprehensive checklist for project handover, code documentation, and academic report submission.",
                        "Coordinated with Kevin to inspect final UI-to-API bindings and ensure all visual assets and architectural copy were finalized across all pages.",
                        "Reviewed the remaining deliverables outlined in the BIT320 handbook to verify that all appendices and submission criteria were accounted for."
                    ],
                    [
                        "Project Handover Planning",
                        "Deliverables Verification",
                        "Cross-Functional Final Audit"
                    ]
                ),
                (
                    "Tuesday,\n29 September 2026",
                    [
                        "Researched, designed, and published the final educational Instagram post on mobile phone fun facts, chronicling the dawn of smartphones from the IBM Simon in 1994 to modern 3nm AI system-on-chips.",
                        "Resolved edge-case bug fixes in the contact inquiry pipeline, configuring the MailModule and Nodemailer transporter to ensure automated inquiry notifications and confirmations dispatch reliably without console warnings.",
                        "Assisted Kevin in fine-tuning the admin dashboard catalog table styles, ensuring action buttons and thumbnail images align cleanly across all screen sizes."
                    ],
                    [
                        "Mobile Computing History Synthesis",
                        "Nodemailer SMTP Dispatch Configuration",
                        "Collaborative Frontend-Backend Bug Fixing"
                    ]
                ),
                (
                    "Wednesday,\n30 September 2026",
                    [
                        "Attended the final project handover meeting with client Bli Dek, the supervisor, and Kevin, officially delivering the completed Grahita Design web platform.",
                        "Provided Bli Dek with the administrative credentials, demonstrated the content management workflows, and explained the automatic cube slot synchronization feature.",
                        "Received formal acceptance and enthusiastic appreciation from Bli Dek for delivering a modern, responsive, and 3D-integrated architecture portfolio platform."
                    ],
                    [
                        "Formal Project Handover Execution",
                        "Client System Handover Skills",
                        "Professional Communication & Demeanor"
                    ]
                ),
                (
                    "Thursday,\n1 October 2026",
                    [
                        "Finalized the technical documentation in README.md, including detailed local setup guides, environment variables, API endpoints list, and production deployment instructions.",
                        "Completed the comprehensive BIT320 Industrial Internship Final Report (FINAL_REPORT_BIT320.md), thoroughly detailing project accomplishments, technical challenges, and personal growth.",
                        "Prepared presentation slides for the final internship defense, summarizing the backend architecture, 3D WebGL integration, and key engineering takeaways."
                    ],
                    [
                        "Technical Documentation Finalization",
                        "Comprehensive Report Synthesis",
                        "Presentation Preparation Skills"
                    ]
                ),
                (
                    "Friday,\n2 October 2026",
                    [
                        "Attended the final internship evaluation and exit meeting with the workplace supervisor at Code Cipta to review overall performance and project achievements.",
                        "Received constructive feedback and supervisor evaluation scores on the BIT320 assessment rubrics, along with the official supervisor signature on internship completion documents.",
                        "Wrote a comprehensive personal reflection on the 14-week internship experience, highlighting technical competence gained in NestJS and TypeORM as well as professional teamwork growth.",
                        "Packaged and submitted all final academic deliverables (Final Report, Logbook Appendix B, Monthly Progress Reports, and repository link) to HELP University."
                    ],
                    [
                        "Professional Performance Appraisal",
                        "Reflective Engineering Assessment",
                        "Formal Academic Submission"
                    ]
                )
            ]
        }
    ]

    for w_idx, week in enumerate(weeks_data):
        # Week Header
        wp = doc.add_paragraph()
        wp.paragraph_format.space_before = Pt(14)
        wp.paragraph_format.space_after = Pt(6)
        wp.paragraph_format.keep_with_next = True
        w_run = wp.add_run(week["week_title"])
        w_run.bold = True
        w_run.font.name = "Calibri"
        w_run.font.size = Pt(13)
        w_run.font.color.rgb = RGBColor(20, 20, 20)

        # Create Table: Rows = 1 (header) + len(entries) + 1 (comments box)
        num_rows = 1 + len(week["entries"]) + 1
        table = doc.add_table(rows=num_rows, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table, color="B0C4DE", sz="4")

        # Set column widths
        col_widths = [Inches(1.5), Inches(3.7), Inches(1.8)]
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width

        # Format Header Row
        hdr_cells = table.rows[0].cells
        hdr_titles = ["Date / Day", "Description of Work Done", "New Skills Learnt"]
        for idx, title in enumerate(hdr_titles):
            cell = hdr_cells[idx]
            set_cell_background(cell, "BDD7EE") # Standard light blue from template
            set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(title)
            run.bold = True
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0, 32, 96)

        # Format Daily Entries
        for e_idx, (date_str, desc_list, skills_list) in enumerate(week["entries"]):
            row_cells = table.rows[1 + e_idx].cells
            
            # Cell 0: Date
            c0 = row_cells[0]
            set_cell_margins(c0, top=100, bottom=100, left=140, right=140)
            p0 = c0.paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p0.paragraph_format.space_before = Pt(2)
            p0.paragraph_format.space_after = Pt(2)
            r0 = p0.add_run(date_str)
            r0.bold = True
            r0.font.name = "Calibri"
            r0.font.size = Pt(9.5)
            r0.font.color.rgb = RGBColor(20, 20, 20)

            # Cell 1: Description of Work Done
            c1 = row_cells[1]
            set_cell_margins(c1, top=100, bottom=100, left=140, right=140)
            for b_idx, bullet in enumerate(desc_list):
                p1 = c1.paragraphs[0] if b_idx == 0 else c1.add_paragraph()
                p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p1.paragraph_format.space_before = Pt(1)
                p1.paragraph_format.space_after = Pt(2)
                p1.paragraph_format.left_indent = Inches(0.15)
                r_dot = p1.add_run("• ")
                r_dot.bold = True
                r_dot.font.name = "Calibri"
                r_dot.font.size = Pt(9)
                r_dot.font.color.rgb = RGBColor(0, 32, 96)
                r1 = p1.add_run(bullet)
                r1.font.name = "Calibri"
                r1.font.size = Pt(9)
                r1.font.color.rgb = RGBColor(30, 30, 30)

            # Cell 2: New Skills Learnt
            c2 = row_cells[2]
            set_cell_margins(c2, top=100, bottom=100, left=140, right=140)
            for s_idx, skill in enumerate(skills_list):
                p2 = c2.paragraphs[0] if s_idx == 0 else c2.add_paragraph()
                p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p2.paragraph_format.space_before = Pt(1)
                p2.paragraph_format.space_after = Pt(2)
                p2.paragraph_format.left_indent = Inches(0.12)
                r_dot2 = p2.add_run("• ")
                r_dot2.bold = True
                r_dot2.font.name = "Calibri"
                r_dot2.font.size = Pt(8.5)
                r_dot2.font.color.rgb = RGBColor(0, 32, 96)
                r2 = p2.add_run(skill)
                r2.font.name = "Calibri"
                r2.font.size = Pt(8.5)
                r2.font.color.rgb = RGBColor(40, 40, 40)

        # Format Comments / Signature Box (bottom row merged across 3 columns)
        comm_row = table.rows[-1]
        c_merged = comm_row.cells[0]
        for col_i in range(1, 3):
            c_merged.merge(comm_row.cells[col_i])
        
        set_cell_background(c_merged, "F8FAFC")
        set_cell_margins(c_merged, top=140, bottom=140, left=160, right=160)
        
        cp0 = c_merged.paragraphs[0]
        cp0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cp0.paragraph_format.space_before = Pt(4)
        cp0.paragraph_format.space_after = Pt(24)
        crun0 = cp0.add_run("Comments by Supervisor / Manager:")
        crun0.bold = True
        crun0.font.name = "Calibri"
        crun0.font.size = Pt(9.5)
        crun0.font.color.rgb = RGBColor(60, 60, 60)

        cp1 = c_merged.add_paragraph()
        cp1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cp1.paragraph_format.space_before = Pt(16)
        cp1.paragraph_format.space_after = Pt(4)
        crun1 = cp1.add_run("Signature: _________________________________________")
        crun1.bold = True
        crun1.font.name = "Calibri"
        crun1.font.size = Pt(9.5)
        crun1.font.color.rgb = RGBColor(60, 60, 60)

        # Add page break between weeks (except after Week 14)
        if w_idx < len(weeks_data) - 1:
            doc.add_page_break()

    output_path = "c:\\Users\\LENOVO\\Downloads\\CLONE GRAHITA\\Internship\\logbook_week10-14.docx"
    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

if __name__ == "__main__":
    build_logbook_docx()
